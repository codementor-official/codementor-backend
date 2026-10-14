"""Trích chữ và cắt đoạn — chạy trong tiến trình con cách ly (`extract_isolated`).

Module này CHỈ import thứ nó cần. Trước đây tiến trình con import `app.rag.documents`, kéo theo
cả gói `app.rag` (`openai`, `boto3`, `fastapi`): ~1 giây import cho MỖI tài liệu, lâu hơn chính
việc trích chữ của đa số tệp.
"""

import io
import json
import sys
import zipfile
from functools import lru_cache

import tiktoken

SUPPORTED_TYPES = ["pdf", "docx", "pptx", "txt", "md"]
MAX_FILE_BYTES = 20 * 1024 * 1024


@lru_cache(maxsize=1)
def tokenizer():
    return tiktoken.get_encoding("cl100k_base")


def chunk_pages(pages: list[dict]) -> list[dict]:
    chunks = []
    for page in pages:
        tokens = tokenizer().encode(page["text"].replace("\x00", "").strip(), disallowed_special=())
        for start in range(0, len(tokens), 520):
            text = tokenizer().decode(tokens[start : start + 600]).strip()
            if text:
                chunks.append({"text": text, "page": page["page"], "embedding": []})
            if len(chunks) > 160:
                raise ValueError(
                    "Tài liệu quá dài (tối đa 160 đoạn). Hãy chia thành tài liệu nhỏ hơn."
                )
            if start + 600 >= len(tokens):
                break
    if not chunks:
        raise ValueError("Không trích xuất được văn bản. PDF scan/ảnh cần OCR trước khi tải lên.")
    return chunks


def extract_pages(kind: str, data: bytes) -> list[dict]:
    if kind not in SUPPORTED_TYPES or len(data) > MAX_FILE_BYTES:
        raise ValueError("Dùng PDF có text, DOCX, PPTX, TXT hoặc Markdown, tối đa 20 MB.")
    if kind in ("txt", "md"):
        try:
            return [{"text": data.decode("utf-8-sig"), "page": None}]
        except UnicodeDecodeError:
            raise ValueError("Vui lòng lưu văn bản ở mã hóa UTF-8.") from None
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("PDF có mật khẩu. Hãy tải lên bản đã bỏ mật khẩu.")
        if len(reader.pages) > 100:
            raise ValueError("Tối đa 100 trang PDF; hãy chia nhỏ tài liệu.")
        return [
            {"text": page.extract_text() or "", "page": i + 1}
            for i, page in enumerate(reader.pages)
        ]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if (
            len(archive.infolist()) > 3000
            or sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024
        ):
            raise ValueError("Tệp Office vượt giới hạn kích thước giải nén.")
    if kind == "pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > 100:
            raise ValueError("Tối đa 100 slide; hãy chia nhỏ tài liệu.")

        def shape_text(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text_frame.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)
                if hasattr(shape, "shapes"):
                    yield from shape_text(shape.shapes)

        return [
            {"text": "\n".join(shape_text(slide.shapes)), "page": i + 1}
            for i, slide in enumerate(presentation.slides)
        ]
    from docx import Document

    document = Document(io.BytesIO(data))
    # Preserve table content; many course documents keep definitions/examples in tables.
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
    )
    return [{"text": text, "page": None}]


def main():
    if sys.platform != "win32":
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    try:
        pages = extract_pages(sys.argv[1], sys.stdin.buffer.read(MAX_FILE_BYTES + 1))
        if sum(len(page["text"]) for page in pages) > 1_000_000:
            raise ValueError("Văn bản trích xuất quá lớn; hãy chia nhỏ tài liệu.")
        result = chunk_pages(pages)
    except ValueError as exc:
        result = {"error": str(exc)[:300]}
    except Exception:
        result = {"error": "Không đọc được tài liệu. Kiểm tra định dạng hoặc mật khẩu."}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    main()
