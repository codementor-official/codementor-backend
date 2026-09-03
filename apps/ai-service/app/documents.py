import asyncio
import io
import json
import math
import subprocess
import sys
import zipfile
from functools import lru_cache

import boto3
import tiktoken
from botocore.config import Config
from fastapi import HTTPException

from app.config import Settings

SUPPORTED_TYPES = ["pdf", "docx", "pptx", "txt", "md"]
MAX_FILE_BYTES = 20 * 1024 * 1024


class DocumentStorage:
    """Read the existing Workspace S3 objects; no second upload or storage flow."""

    def __init__(self, config: Settings):
        self.config = config

    def read(self, workspace_id: str, source: dict) -> bytes:
        key = source.get("storageKey")
        if not key:
            if source["docType"] in ("txt", "md") and source.get("previewText"):
                return source["previewText"].encode("utf-8")
            raise HTTPException(422, "Tài liệu chưa có tệp gốc trên storage. Hãy tải lại tệp.")
        expected = (
            "/".join(
                filter(
                    None,
                    [self.config.aws_s3_document_prefix.strip("/"), "workspaces", workspace_id],
                )
            )
            + "/"
        )
        if not key.startswith(expected) or ".." in key:
            raise HTTPException(403, "Tài liệu không thuộc Workspace này.")
        if not self.config.aws_s3_bucket:
            raise HTTPException(503, "Chưa cấu hình kho lưu trữ tài liệu.")
        credentials = {}
        if self.config.aws_access_key_id.get_secret_value():
            credentials = {
                "aws_access_key_id": self.config.aws_access_key_id.get_secret_value(),
                "aws_secret_access_key": self.config.aws_secret_access_key.get_secret_value(),
            }
            if self.config.aws_session_token.get_secret_value():
                credentials["aws_session_token"] = self.config.aws_session_token.get_secret_value()
        # Construct a session per worker read; default credential chain supports instance roles.
        client = boto3.session.Session().client(
            "s3",
            region_name=self.config.aws_region,
            endpoint_url=self.config.aws_s3_endpoint or None,
            config=Config(
                connect_timeout=10,
                read_timeout=20,
                retries={"max_attempts": 1},
                s3={"addressing_style": "path" if self.config.aws_s3_force_path_style else "auto"},
            ),
            **credentials,
        )
        try:
            result = client.get_object(Bucket=self.config.aws_s3_bucket, Key=key)
            body = result["Body"]
            limit = min(MAX_FILE_BYTES, self.config.document_max_upload_mb * 1024 * 1024)
            try:
                if result.get("ContentLength", 0) > limit:
                    raise HTTPException(422, "Tài liệu vượt giới hạn 20 MB.")
                data = body.read(limit + 1)
                if len(data) > limit:
                    raise HTTPException(422, "Tài liệu vượt giới hạn 20 MB.")
                return data
            finally:
                body.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(
                502, "Không đọc được tài liệu từ storage. Kiểm tra tệp và ENV."
            ) from None
        finally:
            client.close()


@lru_cache(maxsize=1)
def tokenizer():
    return tiktoken.get_encoding("cl100k_base")


def chunk_pages(pages: list[dict]) -> list[dict]:
    chunks = []
    for page in pages:
        tokens = tokenizer().encode(page["text"].replace("\x00", "").strip(), disallowed_special=())
        for start in range(0, len(tokens), 520):
            text = tokenizer().decode(tokens[start : start + 600]).strip()
            if text:
                chunks.append({"text": text, "page": page["page"], "embedding": []})
            if len(chunks) > 160:
                raise ValueError(
                    "Tài liệu quá dài (tối đa 160 đoạn). Hãy chia thành tài liệu nhỏ hơn."
                )
            if start + 600 >= len(tokens):
                break
    if not chunks:
        raise ValueError("Không trích xuất được văn bản. PDF scan/ảnh cần OCR trước khi tải lên.")
    return chunks


def extract_pages(kind: str, data: bytes) -> list[dict]:
    if kind not in SUPPORTED_TYPES or len(data) > MAX_FILE_BYTES:
        raise ValueError("Dùng PDF có text, DOCX, PPTX, TXT hoặc Markdown, tối đa 20 MB.")
    if kind in ("txt", "md"):
        try:
            return [{"text": data.decode("utf-8-sig"), "page": None}]
        except UnicodeDecodeError:
            raise ValueError("Vui lòng lưu văn bản ở mã hóa UTF-8.") from None
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("PDF có mật khẩu. Hãy tải lên bản đã bỏ mật khẩu.")
        if len(reader.pages) > 100:
            raise ValueError("Tối đa 100 trang PDF; hãy chia nhỏ tài liệu.")
        return [
            {"text": page.extract_text() or "", "page": i + 1}
            for i, page in enumerate(reader.pages)
        ]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if (
            len(archive.infolist()) > 3000
            or sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024
        ):
            raise ValueError("Tệp Office vượt giới hạn kích thước giải nén.")
    if kind == "pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > 100:
            raise ValueError("Tối đa 100 slide; hãy chia nhỏ tài liệu.")

        def shape_text(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text_frame.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)
                if hasattr(shape, "shapes"):
                    yield from shape_text(shape.shapes)

        return [
            {"text": "\n".join(shape_text(slide.shapes)), "page": i + 1}
            for i, slide in enumerate(presentation.slides)
        ]
    from docx import Document

    document = Document(io.BytesIO(data))
    # Preserve table content; many course documents keep definitions/examples in tables.
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
    )
    return [{"text": text, "page": None}]


async def extract_isolated(kind: str, data: bytes) -> list[dict]:
    """Parser isolation also works on Windows' Uvicorn reload SelectorEventLoop."""
    process = subprocess.Popen(
        [sys.executable, "-m", "app.extract_worker", kind],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
    )
    communication = asyncio.create_task(asyncio.to_thread(process.communicate, data, timeout=40))
    try:
        output, _ = await asyncio.shield(communication)
    except (subprocess.TimeoutExpired, asyncio.CancelledError) as exc:
        if process.poll() is None:
            process.kill()
        # The communicate thread owns the pipes until it has finished. Do not close them early.
        try:
            await asyncio.shield(communication)
        except subprocess.TimeoutExpired:
            await asyncio.to_thread(process.communicate)
        if isinstance(exc, subprocess.TimeoutExpired):
            raise TimeoutError("Document extraction timed out") from None
        raise
    finally:
        for stream in (process.stdin, process.stdout):
            if stream:
                stream.close()
    if process.returncode or len(output) > 2_000_000:
        raise HTTPException(422, "Không thể xử lý tài liệu này. Hãy kiểm tra hoặc chia nhỏ tệp.")
    result = json.loads(output)
    if isinstance(result, dict):
        raise HTTPException(422, result["error"])
    return result


def cosine(left: list[float], right: list[float]) -> float:
    if not left or len(left) != len(right):
        return 0
    denominator = math.sqrt(sum(x * x for x in left) * sum(x * x for x in right))
    return sum(a * b for a, b in zip(left, right, strict=True)) / denominator if denominator else 0
