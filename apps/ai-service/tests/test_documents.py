import asyncio
import io
import sys

import pytest
from docx import Document
from fastapi import HTTPException
from pypdf import PdfWriter

from app.config import Settings
from app.documents import (
    DocumentStorage,
    chunk_pages,
    cosine,
    extract_isolated,
    extract_pages,
    tokenizer,
)


def test_requested_model_defaults_and_no_key():
    config = Settings(_env_file=None)
    assert config.openai_embedding_model == "text-embedding-3-small"
    assert config.openai_chat_model == "gpt-5-nano"
    assert not config.configured


def test_chunk_token_limit_overlap_and_page():
    text = "Stack lưu dữ liệu theo LIFO. " * 150
    chunks = chunk_pages([{"page": 3, "text": text}])
    assert len(chunks) > 1
    assert all(len(tokenizer().encode(c["text"])) <= 602 for c in chunks)
    assert all(c["page"] == 3 for c in chunks)
    assert all(c["embedding"] == [] for c in chunks)


def test_empty_document_is_not_fabricated():
    with pytest.raises(ValueError, match="OCR"):
        chunk_pages([{"page": 1, "text": ""}])


def test_chunk_limit_rejects_instead_of_silently_truncating():
    with pytest.raises(ValueError, match="160"):
        chunk_pages([{"page": None, "text": "Stack " * 100000}])


def test_docx_reads_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Hàng đợi sử dụng FIFO.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "enqueue"
    table.cell(0, 1).text = "Thêm phần tử"
    output = io.BytesIO()
    doc.save(output)
    pages = extract_pages("docx", output.getvalue())
    assert "FIFO" in pages[0]["text"]
    assert "enqueue | Thêm phần tử" in pages[0]["text"]
    assert pages[0]["page"] is None


def test_pdf_scan_returns_no_text_and_encrypted_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    output = io.BytesIO()
    writer.write(output)
    assert extract_pages("pdf", output.getvalue()) == [{"text": "", "page": 1}]
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="mật khẩu"):
        extract_pages("pdf", output.getvalue())


@pytest.mark.parametrize("kind", ["doc", "link", "png"])
def test_unsupported_formats_fail(kind):
    with pytest.raises(ValueError):
        extract_pages(kind, b"example")


async def test_real_extraction_subprocess():
    chunks = await extract_isolated("md", b"# Stack\nLIFO: last in, first out.")
    assert len(chunks) == 1
    assert "LIFO" in chunks[0]["text"]


@pytest.mark.skipif(sys.platform != "win32", reason="Windows reload event loop")
def test_extraction_on_windows_selector_loop():
    loop = asyncio.SelectorEventLoop()
    try:
        chunks = loop.run_until_complete(extract_isolated("txt", b"Dijkstra uses a heap."))
        assert "Dijkstra" in chunks[0]["text"]
    finally:
        loop.run_until_complete(loop.shutdown_default_executor())
        loop.close()


def test_pptx_reads_text_tables_and_preserves_slide_numbers():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Graph BFS"
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Queue"
    table.cell(0, 1).text = "FIFO"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "DFS uses a stack."
    output = io.BytesIO()
    deck.save(output)
    pages = extract_pages("pptx", output.getvalue())
    assert pages[0]["page"] == 1
    assert "Graph BFS" in pages[0]["text"]
    assert "Queue | FIFO" in pages[0]["text"]
    assert pages[1] == {"page": 2, "text": "DFS uses a stack."}


def test_storage_rejects_foreign_key_without_network():
    storage = DocumentStorage(Settings(_env_file=None))
    with pytest.raises(HTTPException) as error:
        storage.read("workspace-a", {"storageKey": "public/workspace-documents/workspaces/b/a.pdf"})
    assert error.value.status_code == 403


def test_no_url_fetch_or_pdf_preview_fallback():
    storage = DocumentStorage(Settings(_env_file=None))
    with pytest.raises(HTTPException):
        storage.read("a", {"docType": "pdf", "storageKey": None, "previewText": "Only metadata"})
    assert (
        storage.read("a", {"docType": "txt", "storageKey": None, "previewText": "Queue"})
        == b"Queue"
    )


def test_cosine():
    assert cosine([1, 0], [1, 0]) == 1
    assert cosine([1, 0], [0, 1]) == 0
    assert cosine([0, 0], [1, 0]) == 0
    assert cosine([1], [1, 0]) == 0
