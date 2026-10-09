"""Dựng bộ tài liệu đánh giá Tutor từ `documents/` vào `build/`, kèm `manifest.json`.

    uv run python eval/tutor/build_documents.py

Nguồn viết bằng Markdown cho dễ sửa; vài tệp được đổi sang DOCX/PPTX để bộ đánh giá đi qua đủ các
trình trích văn bản thật của `app/rag/documents.py`. Id tài liệu cố định theo thứ tự tệp, nên
seed lại và bộ câu hỏi luôn trỏ đúng tài liệu.
"""

import json
import re
import shutil
from pathlib import Path

from docx import Document
from pptx import Presentation

HERE = Path(__file__).parent
SOURCE = HERE / "documents"
BUILD = HERE / "build"
AS_DOCX = {"02-list-va-tuple.md", "08-lap-trinh-huong-doi-tuong.md"}
AS_PPTX = {"06-de-quy.md", "12-do-thi-bfs-dfs-topo.md"}


def doc_id(number: int) -> str:
    return f"e7a10000-0000-4000-8000-{number:012d}"


def sections(markdown: str) -> tuple[str, list[tuple[str, list[str]]]]:
    title, parts, current = "", [], None
    for line in markdown.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
        elif line.startswith("## "):
            current = (line[3:].strip(), [])
            parts.append(current)
        elif current is not None:
            current[1].append(line)
    return title, parts


def to_docx(markdown: str, target: Path) -> None:
    title, parts = sections(markdown)
    document = Document()
    document.add_heading(title, level=0)
    for heading, lines in parts:
        document.add_heading(heading, level=1)
        for block in re.split(r"\n\s*\n", "\n".join(lines)):
            if block.strip():
                document.add_paragraph(block.strip())
    document.save(target)


def to_pptx(markdown: str, target: Path) -> None:
    title, parts = sections(markdown)
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    for heading, lines in parts:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = heading
        slide.placeholders[1].text_frame.text = "\n".join(line for line in lines if line.strip())
    deck.save(target)


def main() -> None:
    shutil.rmtree(BUILD, ignore_errors=True)
    BUILD.mkdir()
    manifest = []
    for number, path in enumerate(sorted(SOURCE.iterdir()), start=1):
        text = path.read_text(encoding="utf-8")
        title = text.splitlines()[0].lstrip("# ").strip().title() if path.suffix == ".txt" else sections(text)[0]
        if path.name in AS_DOCX:
            target = BUILD / f"{path.stem}.docx"
            to_docx(text, target)
        elif path.name in AS_PPTX:
            target = BUILD / f"{path.stem}.pptx"
            to_pptx(text, target)
        else:
            target = BUILD / path.name
            shutil.copy(path, target)
        manifest.append(
            {
                "id": doc_id(number),
                "key": path.stem,
                "title": title,
                "file": target.name,
                "docType": target.suffix[1:].upper(),
                "sizeBytes": target.stat().st_size,
            }
        )
    (BUILD / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    for item in manifest:
        print(item["id"], item["docType"].ljust(4), item["title"])


if __name__ == "__main__":
    main()
